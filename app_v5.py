"""
AI Job Hunt Hub — Test Build v0.5
Streamlit app that tracks job applications and uses AI (Groq + Sea Lion)
to tailor resumes and generate cover letters.

Now with:
- Supabase (PostgreSQL) persistent storage
- Human-in-the-loop editing for resume & cover letter
- PDF and DOCX export
- LLM Monitoring dashboard

Run:
    streamlit run app.py
"""

import os
import time
import io
from datetime import datetime

import pandas as pd
import streamlit as st
from dotenv import load_dotenv
from sqlalchemy import Column, DateTime, Float, Integer, String, Text, create_engine
from sqlalchemy.orm import declarative_base, sessionmaker
from sqlalchemy.pool import NullPool
from groq import Groq
from openai import OpenAI

# Document generation
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from docx import Document
from docx.shared import Inches

# Rich text editor
from streamlit_lexical_extended import streamlit_lexical_extended

# ---------------------------------------------------------------------------
# 1. CONFIGURATION
# ---------------------------------------------------------------------------

st.set_page_config(
    page_title="Job Hunt Hub",
    page_icon="💼",
    layout="wide",
    initial_sidebar_state="expanded",
)

load_dotenv()

# ============================================================
# 🚀 API KEY CONFIGURATION
# ============================================================

def _get_secret(key: str, default=None):
    try:
        from google.colab import userdata
        val = userdata.get(key)
        if val:
            return val
    except ImportError:
        pass

    val = os.environ.get(key)
    if val:
        return val

    try:
        val = st.secrets.get(key)
        if val:
            return val
    except Exception:
        pass

    return default

GROQ_API_KEY    = _get_secret("GROQ_API_KEY")
SEALION_API_KEY = _get_secret("SEALION_API_KEY")
HF_TOKEN        = _get_secret("HF_TOKEN")

# ============================================================
# 🗄️ DATABASE CONFIGURATION (SUPABASE)
# ============================================================

# Get the connection string from Streamlit secrets.
# The string must start with "postgresql+psycopg2://" for SQLAlchemy.
DATABASE_URL = _get_secret("DATABASE_URL")

if not DATABASE_URL:
    st.error(
        "❌ DATABASE_URL is not set. "
        "Please add your Supabase connection string to `.streamlit/secrets.toml`."
    )
    st.stop()

# ============================================================
# 🚀 API KEY VALIDATION
# ============================================================

if not GROQ_API_KEY and not SEALION_API_KEY:
    st.error(
        "❌ No API keys found. Set at least one of GROQ_API_KEY or SEALION_API_KEY "
        "in your `.env` file or Streamlit secrets."
    )
    st.stop()

if HF_TOKEN:
    os.environ["HF_TOKEN"] = HF_TOKEN

groq_client = Groq(api_key=GROQ_API_KEY) if GROQ_API_KEY else None

SEALION_BASE_URL = "https://api.sea-lion.ai/v1"
sealion_client = (
    OpenAI(api_key=SEALION_API_KEY, base_url=SEALION_BASE_URL)
    if SEALION_API_KEY else None
)

# Constants -----------------------------------------------------------------

STAGES = [
    {"id": "Applied",      "name": "Applied",      "color": "#3498db"},
    {"id": "Acknowledged", "name": "Acknowledged", "color": "#f39c12"},
    {"id": "Screening",    "name": "Screening",    "color": "#9b59b6"},
    {"id": "Interviewing", "name": "Interviewing", "color": "#e67e22"},
    {"id": "Offer",        "name": "Offer",        "color": "#27ae60"},
    {"id": "Rejected",     "name": "Rejected",     "color": "#e74c3c"},
]

# ============================================================
# 🧠 MODEL REGISTRY
# ============================================================

MODEL_REGISTRY = {
    "Sea-Lion v4 27B": {
        "provider": "sealion",
        "model_id": "aisingapore/Gemma-SEA-LION-v4-27B-IT",
        "description": "Gemma 3 27B fine-tuned for SEA languages by AI Singapore.",
    },
    "Sea-Lion v4 32B (Qwen)": {
        "provider": "sealion",
        "model_id": "aisingapore/Qwen-SEA-LION-v4-32B-IT",
        "description": "Qwen3-32B fine-tuned for SEA tasks.",
    },
    "Groq: GPT-OSS 120B": {
        "provider": "groq",
        "model_id": "openai/gpt-oss-120b",
        "description": "OpenAI's open-weight MoE, 120B params, reasoning-capable.",
    },
    "Groq: GPT-OSS 20B": {
        "provider": "groq",
        "model_id": "openai/gpt-oss-20b",
        "description": "Smaller, faster OpenAI open-weight model.",
    },
}

DEFAULT_MODEL_KEY = "Sea-Lion v4 27B"

# ============================================================
# 💰 MODEL PRICING (per 1M tokens, USD)
# ============================================================

MODEL_PRICING = {
    # --- Groq: GPT-OSS 120B ---
    "openai/gpt-oss-120b": {
        "input": 0.15,
        "cached_input": 0.075,
        "output": 0.60,
        "context_window": 131072,
        "max_output": 65536,
    },
    # --- Groq: GPT-OSS 20B ---
    "openai/gpt-oss-20b": {
        "input": 0.075,
        "cached_input": 0.037,
        "output": 0.30,
        "context_window": 131072,
        "max_output": 65536,
    },
    # --- AI Singapore Sea-Lion (free tier) ---
    "aisingapore/Gemma-SEA-LION-v4-27B-IT": {
        "input": 0.0,
        "cached_input": 0.0,
        "output": 0.0,
        "context_window": 131072,
        "max_output": 65536,
    },
    "aisingapore/Qwen-SEA-LION-v4-32B-IT": {
        "input": 0.0,
        "cached_input": 0.0,
        "output": 0.0,
        "context_window": 131072,
        "max_output": 65536,
    },
}

def get_pricing(model_id: str) -> dict:
    return MODEL_PRICING.get(
        model_id,
        {"input": 0.0, "output": 0.0, "cached_input": 0.0,
         "context_window": 0, "max_output": 0},
    )

def compute_cost_usd(model_id: str, input_tokens: int, output_tokens: int) -> float:
    p = get_pricing(model_id)
    return (
        (input_tokens / 1_000_000) * p["input"]
        + (output_tokens / 1_000_000) * p["output"]
    )

# ---------------------------------------------------------------------------
# 2. DATABASE LAYER
# ---------------------------------------------------------------------------

Base = declarative_base()

class Application(Base):
    __tablename__ = "applications"

    id = Column(Integer, primary_key=True)
    company = Column(String, nullable=False)
    role = Column(String, nullable=False)
    job_url = Column(String, default="")
    status = Column(String, default="Applied")
    job_description = Column(Text, default="")
    resume_text = Column(Text, default="")
    cover_letter = Column(Text, default="")
    match_score = Column(String, default="")
    applied_date = Column(DateTime, default=datetime.now)
    last_updated = Column(DateTime, default=datetime.now, onupdate=datetime.now)

    def to_dict(self):
        return {c.name: getattr(self, c.name) for c in self.__table__.columns}

class LLMUsage(Base):
    __tablename__ = "llm_usage"

    id = Column(Integer, primary_key=True)
    called_at = Column(DateTime, default=datetime.now, index=True)
    provider = Column(String)
    model_key = Column(String)
    model_id = Column(String)
    operation = Column(String)
    input_tokens = Column(Integer, default=0)
    output_tokens = Column(Integer, default=0)
    total_tokens = Column(Integer, default=0)
    latency_seconds = Column(Float, default=0.0)
    cost_usd = Column(Float, default=0.0)
    success = Column(Integer, default=1)
    error_message = Column(Text, default="")

    def to_dict(self):
        return {c.name: getattr(self, c.name) for c in self.__table__.columns}

@st.cache_resource
def get_engine():
    # Use NullPool for serverless environments like Streamlit Cloud.
    # The connection string must use "postgresql+psycopg2://" and include "?sslmode=require".
    engine = create_engine(DATABASE_URL, poolclass=NullPool)
    return engine

@st.cache_resource
def get_session_factory():
    return sessionmaker(bind=get_engine(), autoflush=False, autocommit=False)

def db_session():
    return get_session_factory()()

# ---------------------------------------------------------------------------
# 3. AI LAYER
# ---------------------------------------------------------------------------

@st.cache_resource(show_spinner="Loading embedding model (first run only)...")
def load_embedding_model():
    try:
        from sentence_transformers import SentenceTransformer
        return SentenceTransformer("all-MiniLM-L6-v2")
    except Exception as e:
        st.warning(f"Embedding model unavailable: {e}")
        return None

def get_active_model() -> tuple[str, str, str]:
    key = st.session_state.get("selected_model_key", DEFAULT_MODEL_KEY)
    entry = MODEL_REGISTRY.get(key, MODEL_REGISTRY[DEFAULT_MODEL_KEY])
    return entry["provider"], entry["model_id"], key

def load_llm_for(provider: str, model_id: str):
    if provider == "groq":
        if not GROQ_API_KEY:
            return None, "Groq API key not set"
        try:
            from langchain_groq import ChatGroq
            llm = ChatGroq(model=model_id, temperature=0.7, api_key=GROQ_API_KEY)
            return llm, f"Groq ({model_id})"
        except Exception as e:
            return None, f"Groq init failed: {e}"

    if provider == "sealion":
        if not SEALION_API_KEY:
            return None, "Sea Lion API key not set"
        try:
            from langchain_openai import ChatOpenAI
            llm = ChatOpenAI(
                model=model_id, temperature=0.7,
                api_key=SEALION_API_KEY, base_url=SEALION_BASE_URL,
            )
            return llm, f"Sea Lion ({model_id})"
        except Exception as e:
            return None, f"Sea Lion init failed: {e}"

    return None, f"Unknown provider: {provider}"

def _extract_usage_from_langchain(result) -> tuple[int, int]:
    um = getattr(result, "usage_metadata", None)
    if um:
        return int(um.get("input_tokens", 0)), int(um.get("output_tokens", 0))
    rm = getattr(result, "response_metadata", None) or {}
    tu = rm.get("token_usage") or rm.get("usage") or {}
    return int(tu.get("prompt_tokens", 0)), int(tu.get("completion_tokens", 0))

def svc_log_llm_usage(
    *, provider, model_key, model_id, operation,
    input_tokens, output_tokens, latency_seconds,
    success=True, error_message="",
):
    input_tokens = int(input_tokens or 0)
    output_tokens = int(output_tokens or 0)
    cost = compute_cost_usd(model_id, input_tokens, output_tokens) if success else 0.0

    with db_session() as db:
        row = LLMUsage(
            provider=provider, model_key=model_key, model_id=model_id,
            operation=operation, input_tokens=input_tokens, output_tokens=output_tokens,
            total_tokens=input_tokens + output_tokens,
            latency_seconds=float(latency_seconds or 0.0),
            cost_usd=cost, success=1 if success else 0,
            error_message=error_message or "",
        )
        db.add(row)
        db.commit()

def calculate_match_score(resume_text: str, job_description: str) -> float:
    if not resume_text or not job_description:
        return 0.0
    model = load_embedding_model()
    if model is None:
        return 0.0
    from sentence_transformers import util
    emb = model.encode([job_description, resume_text], convert_to_tensor=True)
    score = util.cos_sim(emb[0], emb[1])
    return float(score[0][0])

def call_llm(prompt: str, operation: str = "generic") -> str:
    provider, model_id, label = get_active_model()
    llm, status = load_llm_for(provider, model_id)
    if llm is None:
        svc_log_llm_usage(
            provider=provider, model_key=label, model_id=model_id,
            operation=operation, input_tokens=0, output_tokens=0,
            latency_seconds=0.0, success=False, error_message=str(status),
        )
        return f"⚠️ {status}"
    start = time.time()
    try:
        result = llm.invoke(prompt)
        latency = time.time() - start
        in_tok, out_tok = _extract_usage_from_langchain(result)
        svc_log_llm_usage(
            provider=provider, model_key=label, model_id=model_id,
            operation=operation, input_tokens=in_tok, output_tokens=out_tok,
            latency_seconds=latency, success=True,
        )
        return getattr(result, "content", str(result))
    except Exception as e:
        latency = time.time() - start
        svc_log_llm_usage(
            provider=provider, model_key=label, model_id=model_id,
            operation=operation, input_tokens=0, output_tokens=0,
            latency_seconds=latency, success=False, error_message=str(e),
        )
        return f"⚠️ LLM call failed ({status}): {e}"

def call_groq_direct(prompt: str, operation: str = "direct_test") -> str:
    if groq_client is None:
        return "⚠️ Groq client not initialized (no API key)."
    provider, model_id, label = get_active_model()
    if provider != "groq":
        return "⚠️ Currently selected model is not a Groq model."
    start = time.time()
    try:
        resp = groq_client.chat.completions.create(
            model=model_id,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.7,
        )
        latency = time.time() - start
        usage = getattr(resp, "usage", None)
        in_tok  = int(getattr(usage, "prompt_tokens", 0) or 0)
        out_tok = int(getattr(usage, "completion_tokens", 0) or 0)
        svc_log_llm_usage(
            provider="groq", model_key=label, model_id=model_id,
            operation=operation, input_tokens=in_tok, output_tokens=out_tok,
            latency_seconds=latency, success=True,
        )
        return resp.choices[0].message.content
    except Exception as e:
        latency = time.time() - start
        svc_log_llm_usage(
            provider="groq", model_key=label, model_id=model_id,
            operation=operation, input_tokens=0, output_tokens=0,
            latency_seconds=latency, success=False, error_message=str(e),
        )
        return f"⚠️ Groq direct call failed ({model_id}): {e}"

def call_sealion_direct(prompt: str, operation: str = "direct_test") -> str:
    if sealion_client is None:
        return "⚠️ Sea Lion client not initialized (no API key)."
    provider, model_id, label = get_active_model()
    if provider != "sealion":
        return "⚠️ Currently selected model is not a Sea Lion model."
    start = time.time()
    try:
        resp = sealion_client.chat.completions.create(
            model=model_id,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.7,
        )
        latency = time.time() - start
        usage = getattr(resp, "usage", None)
        in_tok  = int(getattr(usage, "prompt_tokens", 0) or 0)
        out_tok = int(getattr(usage, "completion_tokens", 0) or 0)
        svc_log_llm_usage(
            provider="sealion", model_key=label, model_id=model_id,
            operation=operation, input_tokens=in_tok, output_tokens=out_tok,
            latency_seconds=latency, success=True,
        )
        return resp.choices[0].message.content
    except Exception as e:
        latency = time.time() - start
        svc_log_llm_usage(
            provider="sealion", model_key=label, model_id=model_id,
            operation=operation, input_tokens=0, output_tokens=0,
            latency_seconds=latency, success=False, error_message=str(e),
        )
        return f"⚠️ Sea Lion direct call failed ({model_id}): {e}"

def tailor_resume(resume_text: str, job_description: str) -> str:
    prompt = f"""You are an expert resume writer.
Rewrite the resume below to emphasize experience relevant to the job description.
Use quantifiable achievements and keywords from the job posting.
Keep the format clean and professional. Do NOT invent experience.

JOB DESCRIPTION:
{job_description}

CURRENT RESUME:
{resume_text}

TAILORED RESUME:"""
    return call_llm(prompt, operation="tailor_resume")

def generate_cover_letter(resume_text: str, job_description: str, company: str, role: str) -> str:
    prompt = f"""Write a professional, concise cover letter for the {role} role at {company}.
Ground it in the resume and job description below. Do NOT use placeholders like [Your Name].

JOB DESCRIPTION:
{job_description}

RESUME:
{resume_text}

COVER LETTER:"""
    return call_llm(prompt, operation="cover_letter")

def extract_pdf_text(uploaded_file) -> str:
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(uploaded_file)
        return "".join((page.extract_text() or "") for page in reader.pages)
    except Exception as e:
        st.error(f"Failed to read PDF: {e}")
        return ""

def classify_email_status(body: str) -> str:
    b = body.lower()
    if any(k in b for k in ["unfortunately", "not moving forward", "other candidates", "rejected"]):
        return "Rejected"
    if any(k in b for k in ["interview", "schedule a call", "speak with"]):
        return "Interviewing"
    if any(k in b for k in ["offer", "pleased to offer"]):
        return "Offer"
    if any(k in b for k in ["received your application", "thank you for applying", "acknowledge"]):
        return "Acknowledged"
    return "Applied"

# ---------------------------------------------------------------------------
# 3b. DOCUMENT EXPORT HELPERS
# ---------------------------------------------------------------------------

def generate_pdf_bytes(text: str, filename: str) -> bytes:
    """Generates a simple PDF from a string of text."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER)
    styles = getSampleStyleSheet()
    story = [Paragraph(text, styles["Normal"])]
    doc.build(story)
    return buffer.getvalue()

def generate_docx_bytes(text: str) -> bytes:
    """Generates a DOCX document from a string of text."""
    doc = Document()
    for line in text.split('\n'):
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

# ---------------------------------------------------------------------------
# 4. SERVICE LAYER
# ---------------------------------------------------------------------------

def svc_list_applications():
    with db_session() as db:
        return [a.to_dict() for a in db.query(Application).order_by(Application.last_updated.desc()).all()]

def svc_get_application(app_id: int):
    with db_session() as db:
        app = db.query(Application).filter(Application.id == app_id).first()
        return app.to_dict() if app else None

def svc_create_application(**kwargs):
    with db_session() as db:
        app = Application(**kwargs)
        db.add(app)
        db.commit()
        db.refresh(app)
        return app.id

def svc_update_status(app_id: int, new_status: str):
    with db_session() as db:
        app = db.query(Application).filter(Application.id == app_id).first()
        if app:
            app.status = new_status
            app.last_updated = datetime.now()
            db.commit()
            return True
    return False

def svc_update_application_text(app_id: int, resume_text: str, cover_letter: str):
    with db_session() as db:
        app = db.query(Application).filter(Application.id == app_id).first()
        if app:
            app.resume_text = resume_text
            app.cover_letter = cover_letter
            app.last_updated = datetime.now()
            db.commit()
            return True
    return False

def svc_delete_application(app_id: int):
    with db_session() as db:
        app = db.query(Application).filter(Application.id == app_id).first()
        if app:
            db.delete(app)
            db.commit()

def svc_list_llm_usage(days: int = 7):
    from datetime import timedelta
    cutoff = datetime.now() - timedelta(days=days)
    with db_session() as db:
        rows = (
            db.query(LLMUsage)
            .filter(LLMUsage.called_at >= cutoff)
            .order_by(LLMUsage.called_at.desc())
            .all()
        )
        return [r.to_dict() for r in rows]

# ---------------------------------------------------------------------------
# 5. UI LAYER
# ---------------------------------------------------------------------------

def render_model_selector():
    st.sidebar.markdown("### 🧠 AI Model")
    options = list(MODEL_REGISTRY.keys())
    default_idx = options.index(DEFAULT_MODEL_KEY) if DEFAULT_MODEL_KEY in options else 0
    selected = st.sidebar.selectbox(
        "Select model", options,
        index=st.session_state.get("selected_model_index", default_idx),
        key="selected_model_widget",
        label_visibility="collapsed",
    )
    st.session_state["selected_model_key"] = selected
    st.session_state["selected_model_index"] = options.index(selected)
    entry = MODEL_REGISTRY[selected]
    provider = entry["provider"]
    st.sidebar.caption(f"Provider: **{provider}**")
    st.sidebar.caption(f"`{entry['model_id']}`")
    st.sidebar.caption(entry["description"])
    if provider == "groq" and not GROQ_API_KEY:
        st.sidebar.error("Groq key missing.")
    if provider == "sealion" and not SEALION_API_KEY:
        st.sidebar.error("Sea Lion key missing.")

def render_sidebar():
    with st.sidebar:
        st.title("💼 AI Job Hunt Hub")
        st.caption("AI-powered job search tracker")
        st.markdown("---")
        st.markdown("**👤 Lester**")
        st.caption("lesterchia1@hotmail.com")
        st.markdown("---")
        nav = st.radio(
            "Navigation",
            ["📊 Tracker", "➕ New Application", "📝 Edit & Export",
             "📥 Outbox", "📈 Analytics", "🔍 Monitoring", "⚙️ Settings"],
            index=0,
            label_visibility="collapsed",
        )
        st.markdown("---")
        search = st.text_input("🔍 Search company / role", "")
        st.markdown("---")
        render_model_selector()
        st.markdown("---")
        provider, model_id, label = get_active_model()
        llm, status = load_llm_for(provider, model_id)
        if llm is None:
            st.error(f"LLM unavailable: {status}")
        else:
            st.success(f"LLM ready: {label}")
    return nav, search

def render_kpi_banner(df: pd.DataFrame):
    total = len(df)
    applied = int((df["status"] == "Applied").sum()) if total else 0
    interviewing = int((df["status"] == "Interviewing").sum()) if total else 0
    offers = int((df["status"] == "Offer").sum()) if total else 0
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Total Applications", total)
    c2.metric("Applied", applied)
    c3.metric("Interviewing", interviewing)
    c4.metric("Offers", offers)

def render_kanban(df: pd.DataFrame):
    cols = st.columns(len(STAGES))
    for col, stage in zip(cols, STAGES):
        stage_df = df[df["status"] == stage["id"]]
        with col:
            st.markdown(
                f"""<div style="background:{stage['color']};color:white;
                padding:6px 10px;border-radius:6px;font-weight:600;
                font-size:13px;text-align:center;">
                {stage['name']} ({len(stage_df)})</div>""",
                unsafe_allow_html=True,
            )
            st.write("")
            if stage_df.empty:
                st.caption("—")
                continue
            for _, row in stage_df.iterrows():
                with st.container(border=True):
                    st.markdown(f"**{row['role']}**")
                    st.caption(f"🏢 {row['company']}")
                    if row.get("match_score"):
                        st.caption(f"🎯 {row['match_score']}")
                    new_status = st.selectbox(
                        "Move to",
                        [s["id"] for s in STAGES],
                        index=[s["id"] for s in STAGES].index(row["status"]),
                        key=f"status_{row['id']}",
                        label_visibility="collapsed",
                    )
                    if new_status != row["status"]:
                        svc_update_status(int(row["id"]), new_status)
                        st.rerun()
                    with st.expander("Details"):
                        st.write(f"**Applied:** {row['applied_date']}")
                        if row.get("job_url"):
                            st.markdown(f"[Job posting]({row['job_url']})")
                        if st.button("🗑 Delete", key=f"del_{row['id']}"):
                            svc_delete_application(int(row["id"]))
                            st.rerun()

def view_tracker(search: str):
    st.title("📊 Tracker")
    st.caption("Track your applications from wishlist to offer.")
    apps = svc_list_applications()
    df = pd.DataFrame(apps) if apps else pd.DataFrame(
        columns=["id", "company", "role", "status", "match_score", "applied_date"]
    )
    if search and not df.empty:
        mask = (
            df["company"].str.contains(search, case=False, na=False)
            | df["role"].str.contains(search, case=False, na=False)
        )
        df = df[mask]
    render_kpi_banner(df)
    st.markdown("---")
    if df.empty:
        st.info("No applications yet. Use **➕ New Application** in the sidebar to add one.")
    else:
        render_kanban(df)

def view_new_application():
    st.title("➕ New Application")
    st.caption("Paste the job description and upload your resume. AI will score, tailor, and draft.")
    _, _, model_label = get_active_model()
    st.info(f"Using model: **{model_label}**")
    with st.form("new_app_form", clear_on_submit=True):
        c1, c2 = st.columns(2)
        company = c1.text_input("Company Name *")
        role = c2.text_input("Role Title *")
        job_url = st.text_input("Job URL (optional)")
        job_description = st.text_area("Job Description", height=180)
        resume_file = st.file_uploader("Upload Resume (PDF)", type=["pdf"])
        submitted = st.form_submit_button("🔍 Analyze & Add", type="primary")
    if submitted:
        if not company or not role:
            st.error("Company and Role are required.")
            return
        resume_text = extract_pdf_text(resume_file) if resume_file else ""
        match_str = ""
        if resume_text and job_description:
            with st.spinner("Scoring resume against job..."):
                score = calculate_match_score(resume_text, job_description)
                match_str = f"{score:.0%}"
                st.metric("Profile Match Score", match_str)
            with st.spinner(f"Tailoring resume with {model_label}..."):
                tailored = tailor_resume(resume_text, job_description)
            with st.spinner(f"Drafting cover letter with {model_label}..."):
                cover = generate_cover_letter(tailored, job_description, company, role)
            st.success("Documents generated.")
            with st.expander("📄 Tailored Resume", expanded=False):
                st.text_area("Tailored Resume", tailored, height=280, label_visibility="collapsed")
            with st.expander("✉️ Cover Letter", expanded=False):
                st.text_area("Cover Letter", cover, height=280, label_visibility="collapsed")
            app_id = svc_create_application(
                company=company, role=role, job_url=job_url,
                job_description=job_description,
                resume_text=tailored, cover_letter=cover,
                match_score=match_str, status="Applied",
            )
            st.success(f"✅ Application #{app_id} saved to Tracker.")
        else:
            app_id = svc_create_application(
                company=company, role=role, job_url=job_url,
                job_description=job_description, status="Applied",
            )
            st.info(
                f"Application #{app_id} saved without AI analysis. "
                "Upload a PDF + paste a job description to enable AI features."
            )

def view_edit_export():
    st.title("📝 Edit & Export Documents")
    st.caption("Refine your AI-generated resume and cover letter, then export as PDF or DOCX.")

    apps = svc_list_applications()
    if not apps:
        st.info("No applications found. Add one first from **➕ New Application**.")
        return

    app_options = {f"#{a['id']} · {a['company']} · {a['role']}": a['id'] for a in apps}
    selected_label = st.selectbox("Select an application to edit", list(app_options.keys()))
    app_id = app_options[selected_label]

    app = svc_get_application(app_id)
    if not app:
        st.error("Application not found.")
        return

    st.markdown(f"**{app['role']}** at **{app['company']}**")
    st.caption(f"Status: {app['status']} · Match: {app.get('match_score', 'N/A')}")
    st.markdown("---")

    tab1, tab2 = st.tabs(["📄 Tailored Resume", "✉️ Cover Letter"])

    with tab1:
        edited_resume = streamlit_lexical_extended(
            value=app.get("resume_text", ""),
            key=f"resume_editor_{app_id}",
            height=500,
            toolbar=["undo", "redo", "block_type", "bold", "italic", "underline", "bullet_list", "numbered_list"],
        )
        col1, col2 = st.columns(2)
        with col1:
            if st.button("💾 Save Resume", key=f"save_resume_{app_id}"):
                svc_update_application_text(app_id, edited_resume, app.get("cover_letter", ""))
                st.success("Resume saved.")
                st.rerun()
        with col2:
            pdf_bytes = generate_pdf_bytes(edited_resume, "resume.pdf")
            st.download_button(
                "📥 Download as PDF",
                data=pdf_bytes,
                file_name=f"{app['company']}_Resume.pdf",
                mime="application/pdf",
                key=f"pdf_resume_{app_id}",
            )

    with tab2:
        edited_cover = streamlit_lexical_extended(
            value=app.get("cover_letter", ""),
            key=f"cover_editor_{app_id}",
            height=500,
            toolbar=["undo", "redo", "block_type", "bold", "italic", "underline", "bullet_list", "numbered_list"],
        )
        col1, col2 = st.columns(2)
        with col1:
            if st.button("💾 Save Cover Letter", key=f"save_cover_{app_id}"):
                svc_update_application_text(app_id, app.get("resume_text", ""), edited_cover)
                st.success("Cover letter saved.")
                st.rerun()
        with col2:
            docx_bytes = generate_docx_bytes(edited_cover)
            st.download_button(
                "📥 Download as DOCX",
                data=docx_bytes,
                file_name=f"{app['company']}_Cover_Letter.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key=f"docx_cover_{app_id}",
            )

def view_outbox():
    st.title("📥 Outbox")
    st.info("Email sync is not yet wired up. This view will show parsed job-related emails.")
    sample = pd.DataFrame({
        "Date": ["2026-10-08", "2026-10-07", "2026-10-06"],
        "Sender": ["recruiting@stripe.com", "no-reply@google.com", "hr@meta.com"],
        "Subject": ["Interview Invitation", "Application Received", "Update on your application"],
        "Detected Status": ["🟢 Interviewing", "🟡 Acknowledged", "🔴 Rejected"],
    })
    st.dataframe(sample, use_container_width=True, hide_index=True)
    st.markdown("**Test the classifier:**")
    body = st.text_area("Paste an email body:")
    if body:
        st.success(f"Detected status: **{classify_email_status(body)}**")

def view_analytics():
    st.title("📈 Analytics")
    apps = svc_list_applications()
    if not apps:
        st.info("No data yet. Add applications to see analytics.")
        return
    df = pd.DataFrame(apps)
    c1, c2 = st.columns(2)
    with c1:
        st.subheader("Applications by Stage")
        st.bar_chart(df["status"].value_counts())
    with c2:
        st.subheader("Applications Over Time")
        if "applied_date" in df:
            df["date"] = pd.to_datetime(df["applied_date"]).dt.date
            st.line_chart(df.groupby("date").size())

# ---------------------------------------------------------------------------
# 5b. MONITORING VIEW
# ---------------------------------------------------------------------------

def _fmt_int(n):
    return f"{int(n):,}"

def _fmt_usd(v):
    return f"${v:,.4f}" if v < 1 else f"${v:,.2f}"

def view_monitoring():
    st.title("🔍 LLM Monitoring")
    st.caption(
        "Every LLM request: tokens, cost, latency, failures. "
        "Recommendations never change anything automatically."
    )
    fc1, fc2, fc3, fc4 = st.columns([2, 2, 2, 2])
    with fc1:
        days = st.selectbox("Window", [1, 7, 14, 30, 90], index=1,
                            format_func=lambda d: f"Last {d} day{'s' if d > 1 else ''}")
    with fc2:
        provider_filter = st.selectbox("Provider", ["All", "groq", "sealion"])
    with fc3:
        model_filter = st.selectbox("Model", ["All"] + list(MODEL_REGISTRY.keys()))
    with fc4:
        status_filter = st.selectbox("Status", ["All", "Success", "Failed"])
    rows = svc_list_llm_usage(days=days)
    if not rows:
        st.info(
            "No LLM calls recorded yet. Generate a resume or cover letter "
            "from **➕ New Application** to populate this dashboard."
        )
        _render_pricing_reference()
        return
    df = pd.DataFrame(rows)
    df["called_at"] = pd.to_datetime(df["called_at"])
    if provider_filter != "All":
        df = df[df["provider"] == provider_filter]
    if model_filter != "All":
        df = df[df["model_key"] == model_filter]
    if status_filter == "Success":
        df = df[df["success"] == 1]
    elif status_filter == "Failed":
        df = df[df["success"] == 0]
    if df.empty:
        st.warning("No data matches the current filters.")
        _render_pricing_reference()
        return
    total_requests = len(df)
    success_df = df[df["success"] == 1]
    failed_df = df[df["success"] == 0]
    total_tokens  = int(df["total_tokens"].sum())
    input_tokens  = int(df["input_tokens"].sum())
    output_tokens = int(df["output_tokens"].sum())
    total_cost    = float(df["cost_usd"].sum())
    avg_latency = float(success_df["latency_seconds"].mean()) if not success_df.empty else 0.0
    p95_latency = float(success_df["latency_seconds"].quantile(0.95)) if not success_df.empty else 0.0
    failure_rate = (len(failed_df) / total_requests * 100) if total_requests else 0.0
    st.markdown("---")
    k1, k2, k3, k4 = st.columns(4)
    k1.metric("LLM Requests", _fmt_int(total_requests), f"{avg_latency:.1f}s avg latency")
    k2.metric("Total Tokens", _fmt_int(total_tokens))
    k3.metric("Input Tokens", _fmt_int(input_tokens))
    k4.metric("Output Tokens", _fmt_int(output_tokens))
    k5, k6, k7, k8 = st.columns(4)
    k5.metric("Estimated Cost", _fmt_usd(total_cost))
    k6.metric("Avg Latency", f"{avg_latency:.1f}s")
    k7.metric("P95 Latency", f"{p95_latency:.1f}s")
    k8.metric("Failed Requests", _fmt_int(len(failed_df)), f"{failure_rate:.1f}% of requests")
    st.markdown("---")
    c_left, c_right = st.columns(2)
    with c_left:
        st.subheader("💰 Cost over time")
        st.caption(f"Estimated USD per day · {days}-day window")
        daily = (
            df.set_index("called_at")
              .resample("D")["cost_usd"].sum()
              .reset_index()
              .rename(columns={"called_at": "date", "cost_usd": "cost"})
        )
        if daily["cost"].sum() == 0:
            st.info("All calls in this window were $0 (free-tier models).")
        else:
            st.line_chart(daily.set_index("date")["cost"])
    with c_right:
        st.subheader("🔢 Token usage over time")
        st.caption("Input vs Output per day")
        daily_tok = (
            df.set_index("called_at")
              .resample("D")[["input_tokens", "output_tokens"]].sum()
              .reset_index()
              .rename(columns={
                  "called_at": "date",
                  "input_tokens": "Input",
                  "output_tokens": "Output",
              })
        )
        st.bar_chart(daily_tok.set_index("date")[["Input", "Output"]])
    st.markdown("---")
    c_left2, c_right2 = st.columns(2)
    with c_left2:
        st.subheader("📊 Usage by model")
        by_model = (
            df.groupby("model_key")
              .agg(
                  Requests=("id", "count"),
                  InputTokens=("input_tokens", "sum"),
                  OutputTokens=("output_tokens", "sum"),
                  TotalTokens=("total_tokens", "sum"),
                  CostUSD=("cost_usd", "sum"),
                  AvgLatency=("latency_seconds", "mean"),
              )
              .reset_index()
              .rename(columns={"model_key": "Model"})
        )
        by_model["AvgLatency"] = by_model["AvgLatency"].round(1)
        by_model["CostUSD"] = by_model["CostUSD"].round(4)
        st.dataframe(by_model, use_container_width=True, hide_index=True)
    with c_right2:
        st.subheader("💵 Cost by model")
        cost_by_model = by_model.set_index("Model")["CostUSD"]
        if cost_by_model.sum() == 0:
            st.info("All selected models are free-tier ($0).")
        else:
            st.bar_chart(cost_by_model)
    st.markdown("---")
    st.subheader("⚙️ Usage by operation")
    by_op = (
        df.groupby("operation")
          .agg(
              Requests=("id", "count"),
              TotalTokens=("total_tokens", "sum"),
              CostUSD=("cost_usd", "sum"),
              AvgLatency=("latency_seconds", "mean"),
          )
          .reset_index()
          .rename(columns={"operation": "Operation"})
    )
    by_op["AvgLatency"] = by_op["AvgLatency"].round(1)
    by_op["CostUSD"] = by_op["CostUSD"].round(4)
    st.dataframe(by_op, use_container_width=True, hide_index=True)
    st.markdown("---")
    if not failed_df.empty:
        st.subheader("❌ Recent failures")
        fail_view = (
            failed_df[["called_at", "model_key", "operation",
                       "latency_seconds", "error_message"]]
            .head(10)
            .rename(columns={
                "called_at": "Time",
                "model_key": "Model",
                "operation": "Op",
                "latency_seconds": "Latency (s)",
                "error_message": "Error",
            })
        )
        st.dataframe(fail_view, use_container_width=True, hide_index=True)
    st.markdown("---")
    with st.expander("📜 Raw call log (last 100)"):
        raw = df.head(100)[[
            "called_at", "provider", "model_key", "operation",
            "input_tokens", "output_tokens", "total_tokens",
            "latency_seconds", "cost_usd", "success",
        ]]
        st.dataframe(raw, use_container_width=True, hide_index=True)
    with st.expander("🗑 Danger zone"):
        if st.button("Clear all LLM usage logs", type="secondary"):
            with db_session() as db:
                db.query(LLMUsage).delete()
                db.commit()
            st.success("LLM usage logs cleared.")
            st.rerun()
    st.markdown("---")
    _render_pricing_reference()

def _render_pricing_reference():
    st.subheader("💲 Model pricing reference")
    st.caption("Per 1M tokens (USD). Groq prices per user's plan; Sea-Lion via AI Singapore is free-tier.")
    pricing_rows = []
    for key, entry in MODEL_REGISTRY.items():
        p = get_pricing(entry["model_id"])
        pricing_rows.append({
            "Model Key": key,
            "Provider": entry["provider"],
            "Model ID": entry["model_id"],
            "Input $/1M": p["input"],
            "Cached Input $/1M": p["cached_input"],
            "Output $/1M": p["output"],
            "Context": p["context_window"],
            "Max Output": p["max_output"],
        })
    st.dataframe(pd.DataFrame(pricing_rows), use_container_width=True, hide_index=True)
    st.info(
        "**Cost formula:**  \n"
        "`Cost = (input_tokens / 1,000,000) × input_price "
        "+ (output_tokens / 1,000,000) × output_price`"
    )

# ---------------------------------------------------------------------------
# 5c. SETTINGS VIEW
# ---------------------------------------------------------------------------

def view_settings():
    st.title("⚙️ Settings")
    provider, model_id, label = get_active_model()
    st.markdown("**Active Model**")
    st.code(
        f"Display key : {label}\n"
        f"Provider    : {provider}\n"
        f"Model ID    : {model_id}",
        language="bash",
    )
    st.markdown("**Database Info**")
    st.write(f"- DATABASE_URL: {'✅ set' if DATABASE_URL else '❌ missing'}")
    if DATABASE_URL:
        # Mask password in display
        masked_url = DATABASE_URL.split("@")[-1] if "@" in DATABASE_URL else DATABASE_URL
        st.code(f"Host: ...@{masked_url}")
    st.markdown("**API Keys Detected**")
    st.write(f"- GROQ_API_KEY: {'✅ set' if GROQ_API_KEY else '❌ missing'}")
    st.write(f"- SEALION_API_KEY: {'✅ set' if SEALION_API_KEY else '❌ missing'}")
    st.write(f"- HF_TOKEN: {'✅ set' if HF_TOKEN else '⚪ not set (optional)'}")
    st.markdown("**Test Direct Calls**")
    if provider == "groq":
        if st.button("Test Groq direct call"):
            with st.spinner(f"Calling Groq ({model_id})..."):
                st.info(call_groq_direct("Reply with: 'Groq OK'"))
    elif provider == "sealion":
        if st.button("Test Sea Lion direct call"):
            with st.spinner(f"Calling Sea Lion ({model_id})..."):
                st.info(call_sealion_direct("Reply with: 'Sea Lion OK'"))
    st.markdown("---")
    st.markdown("**Available Models**")
    for key, entry in MODEL_REGISTRY.items():
        marker = " ← active" if key == label else ""
        st.markdown(f"- **{key}**{marker}  \n  `{entry['model_id']}`  \n  {entry['description']}")
    st.markdown("---")
    st.markdown("**Danger Zone**")
    col1, col2 = st.columns(2)
    with col1:
        if st.button("🗑 Clear all applications", type="secondary"):
            with db_session() as db:
                db.query(Application).delete()
                db.commit()
            st.success("All applications cleared.")
            st.rerun()
    with col2:
        if st.button("🗑 Clear LLM usage logs", type="secondary"):
            with db_session() as db:
                db.query(LLMUsage).delete()
                db.commit()
            st.success("LLM usage logs cleared.")
            st.rerun()

# ---------------------------------------------------------------------------
# 6. MAIN
# ---------------------------------------------------------------------------

def main():
    if "selected_model_key" not in st.session_state:
        st.session_state["selected_model_key"] = DEFAULT_MODEL_KEY
    nav, search = render_sidebar()
    if nav == "📊 Tracker":
        view_tracker(search)
    elif nav == "➕ New Application":
        view_new_application()
    elif nav == "📝 Edit & Export":
        view_edit_export()
    elif nav == "📥 Outbox":
        view_outbox()
    elif nav == "📈 Analytics":
        view_analytics()
    elif nav == "🔍 Monitoring":
        view_monitoring()
    elif nav == "⚙️ Settings":
        view_settings()

if __name__ == "__main__":
    main()
